from __future__ import annotations

import csv
import json
import math
import re
import subprocess
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEFAULT_HOURS_PER_MONTH = 744.0
DEFAULT_UNMAPPED_LIMIT = 200
REQUIRED_BILL_COLUMNS = {
    "aws_product_name1_x",
    "aws_item_description",
    "aws_product_quantity",
}


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text == "0" else text


def _num(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return 0.0
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text or text == "-":
        return 0.0
    text = re.sub(r"^\$+\s*", "", text)
    try:
        return float(text)
    except ValueError:
        return 0.0


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _tokens(text: str) -> set[str]:
    return {part for part in _norm(text).split() if part}


def _round(value: float) -> float:
    if abs(value) < 0.000001:
        return 0.0
    return round(value, 6)


@dataclass(frozen=True)
class LimitService:
    name: str
    description: str | None = None


@dataclass(frozen=True)
class LimitDefinition:
    service_name: str
    name: str
    description: str | None = None
    scope_type: str | None = None
    is_deprecated: bool = False
    is_dynamic: bool = False
    is_eligible_for_limit_increase: bool = False
    is_resource_availability_supported: bool = False
    are_quotas_supported: bool = False


@dataclass(frozen=True)
class LimitMatch:
    service_name: str
    service_description: str | None
    limit_name: str
    limit_description: str | None
    scope_type: str | None
    score: int


class LimitCatalog:
    def __init__(
        self,
        *,
        services: Iterable[LimitService],
        definitions: Iterable[LimitDefinition],
        generated_at: str | None = None,
        region: str | None = None,
        source_commands: list[list[str]] | None = None,
    ) -> None:
        self.generated_at = generated_at
        self.region = region
        self.source_commands = source_commands or []
        self.services = {service.name: service for service in services if service.name}
        self.definitions_by_service: dict[str, list[LimitDefinition]] = defaultdict(list)
        for definition in definitions:
            if definition.service_name and definition.name:
                self.definitions_by_service[definition.service_name].append(definition)

    @classmethod
    def from_dict(cls, payload: dict[str, Any]) -> "LimitCatalog":
        services = [
            LimitService(name=item.get("name", ""), description=item.get("description"))
            for item in payload.get("services", [])
        ]
        definitions = [
            LimitDefinition(
                service_name=item.get("service_name") or item.get("service-name") or "",
                name=item.get("name", ""),
                description=item.get("description"),
                scope_type=item.get("scope_type") or item.get("scope-type"),
                is_deprecated=bool(item.get("is_deprecated", item.get("is-deprecated", False))),
                is_dynamic=bool(item.get("is_dynamic", item.get("is-dynamic", False))),
                is_eligible_for_limit_increase=bool(
                    item.get(
                        "is_eligible_for_limit_increase",
                        item.get("is-eligible-for-limit-increase", False),
                    )
                ),
                is_resource_availability_supported=bool(
                    item.get(
                        "is_resource_availability_supported",
                        item.get("is-resource-availability-supported", False),
                    )
                ),
                are_quotas_supported=bool(
                    item.get("are_quotas_supported", item.get("are-quotas-supported", False))
                ),
            )
            for item in payload.get("definitions", [])
        ]
        return cls(
            services=services,
            definitions=definitions,
            generated_at=payload.get("generated_at"),
            region=payload.get("region"),
            source_commands=payload.get("source_commands", []),
        )

    @classmethod
    def load(cls, path: Path) -> "LimitCatalog":
        return cls.from_dict(json.loads(path.read_text(encoding="utf-8")))

    @classmethod
    def from_oci_cli(
        cls,
        *,
        compartment_id: str,
        region: str,
        profile: str = "DEFAULT",
        cli_path: str = "oci",
        timeout_seconds: int = 90,
    ) -> "LimitCatalog":
        base = [cli_path, "--profile", profile, "--region", region, "--output", "json"]
        service_cmd = base + [
            "limits",
            "service",
            "list",
            "--compartment-id",
            compartment_id,
            "--all",
        ]
        definition_cmd = base + [
            "limits",
            "definition",
            "list",
            "--compartment-id",
            compartment_id,
            "--all",
        ]
        services_payload = _run_json_command(service_cmd, timeout_seconds=timeout_seconds)
        definitions_payload = _run_json_command(definition_cmd, timeout_seconds=timeout_seconds)
        services = [
            LimitService(name=item.get("name", ""), description=item.get("description"))
            for item in services_payload.get("data", [])
            if item.get("name")
        ]
        definitions = [
            LimitDefinition(
                service_name=item.get("service-name", ""),
                name=item.get("name", ""),
                description=item.get("description"),
                scope_type=item.get("scope-type"),
                is_deprecated=bool(item.get("is-deprecated", False)),
                is_dynamic=bool(item.get("is-dynamic", False)),
                is_eligible_for_limit_increase=bool(
                    item.get("is-eligible-for-limit-increase", False)
                ),
                is_resource_availability_supported=bool(
                    item.get("is-resource-availability-supported", False)
                ),
                are_quotas_supported=bool(item.get("are-quotas-supported", False)),
            )
            for item in definitions_payload.get("data", [])
            if item.get("service-name") and item.get("name")
        ]
        return cls(
            services=services,
            definitions=definitions,
            generated_at=datetime.now(UTC).isoformat(),
            region=region,
            source_commands=[service_cmd, definition_cmd],
        )

    def to_dict(self) -> dict[str, Any]:
        definitions = [
            asdict(definition)
            for definition in sorted(
                (d for defs in self.definitions_by_service.values() for d in defs),
                key=lambda d: (d.service_name, d.name),
            )
        ]
        return {
            "generated_at": self.generated_at,
            "region": self.region,
            "source_commands": self.source_commands,
            "services": [
                asdict(service) for service in sorted(self.services.values(), key=lambda s: s.name)
            ],
            "definitions": definitions,
        }

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(self.to_dict(), indent=2, sort_keys=True) + "\n", encoding="utf-8")

    def service_description(self, service_name: str) -> str | None:
        service = self.services.get(service_name)
        return service.description if service else None

    def has_service(self, service_name: str) -> bool:
        return service_name in self.services

    def match_limit(
        self,
        *,
        service_name: str,
        terms: Iterable[str],
        preferred_names: Iterable[str] = (),
        avoid_terms: Iterable[str] = (),
        minimum_score: int = 8,
    ) -> LimitMatch | None:
        if service_name not in self.services:
            return None
        candidates = self.definitions_by_service.get(service_name, [])
        scored: list[tuple[int, LimitDefinition]] = []
        term_list = [term for term in terms if term]
        preferred_list = [name for name in preferred_names if name]
        avoid_list = [term for term in avoid_terms if term]

        for definition in candidates:
            haystack = f"{definition.name} {definition.description or ''}".lower()
            haystack_norm = _norm(haystack)
            haystack_tokens = _tokens(haystack)
            score = 0
            for preferred_name in preferred_list:
                preferred = preferred_name.lower()
                if definition.name == preferred:
                    score += 35
                elif definition.name.startswith(preferred):
                    score += 22
                elif preferred in definition.name:
                    score += 14
            for term in term_list:
                score += _term_score(haystack_norm, haystack_tokens, term)
            for avoid in avoid_list:
                if _term_score(haystack_norm, haystack_tokens, avoid) > 0:
                    score -= 18
            if definition.is_deprecated:
                score -= 30
            if definition.is_resource_availability_supported:
                score += 2
            if definition.are_quotas_supported:
                score += 1
            if definition.is_eligible_for_limit_increase:
                score += 1
            scored.append((score, definition))

        scored.sort(key=lambda item: (item[0], not item[1].is_deprecated, item[1].name), reverse=True)
        if not scored or scored[0][0] < minimum_score:
            return None
        best_score, best = scored[0]
        return LimitMatch(
            service_name=service_name,
            service_description=self.service_description(service_name),
            limit_name=best.name,
            limit_description=best.description,
            scope_type=best.scope_type,
            score=best_score,
        )


def _term_score(haystack_norm: str, haystack_tokens: set[str], term: str) -> int:
    term_norm = _norm(term)
    if not term_norm:
        return 0
    if term_norm in haystack_norm:
        return 10
    term_tokens = set(term_norm.split())
    if term_tokens and term_tokens.issubset(haystack_tokens):
        return 7
    compact = term_norm.replace(" ", "")
    compact_haystack = haystack_norm.replace(" ", "")
    if compact and compact in compact_haystack:
        return 6
    return 0


def _run_json_command(command: list[str], *, timeout_seconds: int) -> dict[str, Any]:
    completed = subprocess.run(
        command,
        check=False,
        capture_output=True,
        text=True,
        timeout=timeout_seconds,
    )
    if completed.returncode != 0:
        message = completed.stderr.strip() or completed.stdout.strip()
        raise RuntimeError(f"Command failed ({completed.returncode}): {' '.join(command)}\n{message}")
    try:
        return json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Command returned invalid JSON: {' '.join(command)}") from exc


@dataclass(frozen=True)
class BillComparisonLine:
    row_number: int
    source_sheet: str
    aws_product: str
    aws_description: str
    quantity: float
    aws_cost: float
    aws_vcpu: float
    aws_memory_gib: float
    aws_nvme_tb: float
    aws_nvme_gb: float
    aws_gpu: float
    oci_service_name: str
    oci_compute_shape: str
    oci_part_num_compute_shape: str
    oci_part_num_memory: str
    oci_part_num_obj_stor: str
    oci_part_num_block_vol: str
    oci_part_num_block_vol_perf: str
    oci_datatransfer_part_num: str
    oci_sku: str
    oci_sku2: str

    @classmethod
    def from_mapping(cls, row_number: int, source_sheet: str, data: dict[str, Any]) -> "BillComparisonLine":
        lookup = {str(key).strip(): value for key, value in data.items()}
        return cls(
            row_number=row_number,
            source_sheet=source_sheet,
            aws_product=_clean(lookup.get("aws_product_name1_x")),
            aws_description=_clean(lookup.get("aws_item_description")),
            quantity=_num(lookup.get("aws_product_quantity")),
            aws_cost=_num(lookup.get("aws_cost")),
            aws_vcpu=_num(lookup.get("aws_vcpu")),
            aws_memory_gib=_num(lookup.get("aws_memory_gib")),
            aws_nvme_tb=_num(lookup.get("aws_NVMe_TB")),
            aws_nvme_gb=_num(lookup.get("aws_NVMe_GB")),
            aws_gpu=_num(lookup.get("aws_gpu")),
            oci_part_num_compute_shape=_clean(lookup.get("oci_part_num_compute_shape")),
            oci_part_num_memory=_clean(lookup.get("oci_part_num_memory")),
            oci_compute_shape=_clean(lookup.get("oci_compute_shape")),
            oci_service_name=_clean(lookup.get("oci_service_name")),
            oci_part_num_obj_stor=_clean(lookup.get("oci_part_num_obj_stor")),
            oci_part_num_block_vol=_clean(lookup.get("oci_part_num_block_vol")),
            oci_part_num_block_vol_perf=_clean(lookup.get("oci_part_num_block_vol_perf")),
            oci_datatransfer_part_num=_clean(lookup.get("oci_datatransfer_part_num")),
            oci_sku=_clean(lookup.get("oci_sku")),
            oci_sku2=_clean(lookup.get("oci_sku2")),
        )

    @property
    def search_text(self) -> str:
        return " ".join(
            [
                self.aws_product,
                self.aws_description,
                self.oci_service_name,
                self.oci_compute_shape,
                self.oci_sku,
                self.oci_sku2,
            ]
        ).lower()

    @property
    def evidence(self) -> str:
        oci = f" -> {self.oci_service_name}" if self.oci_service_name else ""
        return f"row {self.row_number}: {self.aws_product}{oci}: {self.aws_description[:160]}"

    def is_hour_usage(self) -> bool:
        text = self.search_text
        return bool(re.search(r"\b(hours?|hrs?|hourly)\b", text))


@dataclass
class RequirementDraft:
    service_name: str
    metric: str
    quantity: float
    unit: str
    terms: list[str]
    preferred_names: list[str]
    assumptions: list[str]
    confidence: str = "medium"
    avoid_terms: list[str] = field(
        default_factory=lambda: ["reserved", "reservable", "free", "backup"]
    )


@dataclass
class RequirementAggregate:
    service_name: str
    service_description: str | None
    limit_name: str
    limit_description: str | None
    scope_type: str | None
    metric: str
    required_quantity: float
    unit: str
    confidence: str
    source_line_count: int
    source_rows: list[int]
    source_products: list[str]
    source_oci_services: list[str]
    assumptions: list[str]
    evidence: list[str]


class BillComparisonAnalyzer:
    def __init__(
        self,
        catalog: LimitCatalog,
        *,
        hours_per_month: float = DEFAULT_HOURS_PER_MONTH,
        unmapped_limit: int = DEFAULT_UNMAPPED_LIMIT,
    ) -> None:
        self.catalog = catalog
        self.hours_per_month = hours_per_month
        self.unmapped_limit = unmapped_limit

    def analyze_file(self, path: Path) -> dict[str, Any]:
        lines = load_bill_comparison_lines(path)
        return self.analyze_lines(lines, source_file=str(path))

    def analyze_lines(self, lines: list[BillComparisonLine], *, source_file: str | None = None) -> dict[str, Any]:
        service_summary: dict[str, dict[str, Any]] = {}
        aggregates: dict[tuple[str, str, str, str], dict[str, Any]] = {}
        unmapped: list[dict[str, Any]] = []
        ignored_reasons: Counter[str] = Counter()
        unmapped_reasons: Counter[str] = Counter()
        mapped_rows: set[int] = set()

        for line in lines:
            inferred_service = self.infer_service(line)
            if inferred_service:
                summary = service_summary.setdefault(
                    inferred_service,
                    {
                        "service_name": inferred_service,
                        "service_description": self.catalog.service_description(inferred_service),
                        "line_items": 0,
                        "source_quantity": 0.0,
                        "aws_products": Counter(),
                        "oci_service_names": Counter(),
                    },
                )
                summary["line_items"] += 1
                summary["source_quantity"] += line.quantity
                if line.aws_product:
                    summary["aws_products"][line.aws_product] += 1
                if line.oci_service_name:
                    summary["oci_service_names"][line.oci_service_name] += 1

            drafts, reason = self._drafts_for_line(line)
            if not drafts:
                if reason.startswith("ignored:"):
                    ignored_reasons[reason.removeprefix("ignored:")] += 1
                else:
                    unmapped_reasons[reason] += 1
                    self._append_unmapped(unmapped, line, reason)
                continue

            row_had_match = False
            for draft in drafts:
                match = self.catalog.match_limit(
                    service_name=draft.service_name,
                    terms=draft.terms,
                    preferred_names=draft.preferred_names,
                    avoid_terms=draft.avoid_terms,
                )
                if not match:
                    reason = f"no_limit_match:{draft.service_name}:{draft.metric}"
                    unmapped_reasons[reason] += 1
                    self._append_unmapped(unmapped, line, reason)
                    continue

                row_had_match = True
                key = (match.service_name, match.limit_name, draft.metric, draft.unit)
                record = aggregates.setdefault(
                    key,
                    {
                        "service_name": match.service_name,
                        "service_description": match.service_description,
                        "limit_name": match.limit_name,
                        "limit_description": match.limit_description,
                        "scope_type": match.scope_type,
                        "metric": draft.metric,
                        "required_quantity": 0.0,
                        "unit": draft.unit,
                        "confidence_values": [],
                        "source_rows": set(),
                        "source_products": Counter(),
                        "source_oci_services": Counter(),
                        "assumptions": set(),
                        "evidence": [],
                    },
                )
                record["required_quantity"] += draft.quantity
                record["confidence_values"].append(draft.confidence)
                record["source_rows"].add(line.row_number)
                if line.aws_product:
                    record["source_products"][line.aws_product] += 1
                if line.oci_service_name:
                    record["source_oci_services"][line.oci_service_name] += 1
                record["assumptions"].update(draft.assumptions)
                if len(record["evidence"]) < 8:
                    record["evidence"].append(line.evidence)
            if row_had_match:
                mapped_rows.add(line.row_number)

        requirements = [self._aggregate_to_output(record) for record in aggregates.values()]
        requirements.sort(key=lambda item: (item["service_name"], item["limit_name"], item["metric"]))
        service_items = [self._service_summary_to_output(item) for item in service_summary.values()]
        service_items.sort(key=lambda item: (-item["line_items"], item["service_name"]))

        return {
            "source_file": source_file,
            "generated_at": datetime.now(UTC).isoformat(),
            "hours_per_month": self.hours_per_month,
            "catalog_generated_at": self.catalog.generated_at,
            "catalog_region": self.catalog.region,
            "summary": {
                "line_items": len(lines),
                "mapped_line_items": len(mapped_rows),
                "requirement_count": len(requirements),
                "unmapped_line_items": sum(unmapped_reasons.values()),
                "ignored_line_items": sum(ignored_reasons.values()),
                "unmapped_reason_counts": dict(unmapped_reasons.most_common()),
                "ignored_reason_counts": dict(ignored_reasons.most_common()),
            },
            "requirements": requirements,
            "service_summary": service_items,
            "unmapped": unmapped,
        }

    def infer_service(self, line: BillComparisonLine) -> str | None:
        text = line.search_text
        service_aliases = [
            ("network-load-balancer-api", ["network load balancer"]),
            ("load-balancer", ["oci load balancer", "application load balancer", "elastic load balancing"]),
            ("block-storage", ["oci block volume", "block volume", "ebs", "gp2", "gp3"]),
            ("object-storage", ["oci object storage", "simple storage service", "s3"]),
            ("postgresql", ["postgresql service", "aurora postgresql", "rds postgresql"]),
            ("mysql", ["mysql", "heatwave"]),
            ("database", ["database service", "relational database service", "oracle dbaas"]),
            ("compute", ["vm.standard", "vm.optimized", "oci compute", "elastic compute cloud", "ec2"]),
            ("vcn", ["virtual private cloud", "public ipv4", "transitgateway", "nat gateway"]),
            ("filesystem", ["file storage", "efs", "fsx"]),
            ("dns", ["route 53", "dns"]),
            ("api-gateway", ["api gateway"]),
            ("monitoring", ["cloudwatch", "monitoring"]),
            ("logging", ["cloudtrail", "logging"]),
            ("queue", ["simple queue service", "sqs", "queue"]),
            ("notifications", ["simple notification service", "sns", "notification"]),
            ("waf", ["web application firewall", "aws waf", "waf"]),
            ("kms", ["key management service", "kms", "vault"]),
        ]
        for service, aliases in service_aliases:
            if service in self.catalog.services and any(alias in text for alias in aliases):
                return service
        normalized_text = _norm(text)
        for service_name in self.catalog.services:
            normalized_service = _norm(service_name)
            if len(normalized_service) > 4 and normalized_service in normalized_text:
                return service_name
        return None

    def _drafts_for_line(self, line: BillComparisonLine) -> tuple[list[RequirementDraft], str]:
        if line.quantity <= 0:
            return [], "ignored:zero_quantity"

        drafts: list[RequirementDraft] = []
        text = line.search_text

        if self._is_database_line(line):
            drafts.extend(self._database_drafts(line))

        if self._is_compute_line(line):
            drafts.extend(self._compute_drafts(line))

        if self._is_block_storage_line(line):
            drafts.extend(self._block_storage_drafts(line))

        if self._is_object_storage_line(line):
            drafts.extend(self._object_storage_drafts(line))

        if "loadbalancer-hour" in text or "load balancer-hour" in text:
            drafts.extend(self._load_balancer_drafts(line))

        if "public ipv4 address per hour" in text or "public ipv4" in text:
            drafts.extend(
                self._single_count_draft(
                    line=line,
                    service_name="vcn",
                    metric="reserved_public_ip_count",
                    terms=["reserved", "public", "ip", "count"],
                    preferred_names=["reserved-public-ip-count"],
                    unit="public_ip",
                    assumptions=["Public IPv4 address-hours were converted to a steady-state count."],
                )
            )

        if "nat gateway" in text and line.is_hour_usage():
            drafts.extend(
                self._single_count_draft(
                    line=line,
                    service_name="vcn",
                    metric="nat_gateway_count",
                    terms=["nat", "gateway", "count"],
                    preferred_names=["nat-gateway-count"],
                    unit="gateway",
                    assumptions=["NAT gateway-hours were converted to a steady-state count."],
                )
            )

        if "transitgateway-hours" in text or "transit gateway" in text and line.is_hour_usage():
            drafts.extend(
                self._single_count_draft(
                    line=line,
                    service_name="vcn",
                    metric="drg_count",
                    terms=["dynamic", "routing", "gateway", "drg", "count"],
                    preferred_names=["drg-count"],
                    unit="gateway",
                    assumptions=["AWS Transit Gateway hours were mapped to OCI DRG count."],
                    confidence="low",
                )
            )

        if drafts:
            return drafts, "mapped"

        if "data transfer" in text or "ingress" in text or "egress" in text:
            return [], "ignored:usage_charge_without_service_limit"
        if "request" in text or "queries" in text or "api" in text:
            return [], "ignored:request_or_api_usage_without_resource_limit"
        if self.infer_service(line):
            return [], "service_detected_without_resource_quantity"
        return [], "no_supported_service_detected"

    def _is_compute_line(self, line: BillComparisonLine) -> bool:
        text = line.search_text
        return bool(
            line.aws_gpu > 0
            or line.oci_part_num_compute_shape
            or line.oci_part_num_memory
            or "amazon elastic compute cloud" in text
            or "amazon elastic container service" in text
            or "amazon elastic kubernetes service" in text
            or "vm.standard" in text
            or "vm.optimized" in text
            or "oci compute" in text
        )

    def _compute_drafts(self, line: BillComparisonLine) -> list[RequirementDraft]:
        drafts: list[RequirementDraft] = []
        instance_count = self._steady_state_count(line)
        family = _compute_family(line)
        assumptions = [
            f"Usage hours were divided by {self.hours_per_month:g} to estimate steady-state resources."
        ]

        if line.aws_gpu > 0:
            gpu_family = _gpu_family(line) or family or "gpu"
            drafts.append(
                RequirementDraft(
                    service_name="compute",
                    metric="compute_gpu_count",
                    quantity=_round(instance_count * line.aws_gpu),
                    unit="gpu",
                    terms=[gpu_family, "gpu", "count"],
                    preferred_names=[f"{gpu_family}-count"] if gpu_family.startswith("gpu-") else [],
                    assumptions=assumptions + ["GPU count was derived from AWS GPU quantity per row."],
                    confidence="high" if gpu_family.startswith("gpu-") else "medium",
                )
            )
            return drafts

        if line.aws_vcpu > 0:
            terms = [family, "core", "count"] if family else ["core", "count"]
            preferred = [f"{family}-core-count"] if family else []
            drafts.append(
                RequirementDraft(
                    service_name="compute",
                    metric="compute_ocpu",
                    quantity=_round(instance_count * line.aws_vcpu / 2.0),
                    unit="ocpu",
                    terms=terms,
                    preferred_names=preferred,
                    assumptions=assumptions + ["AWS vCPU was converted to OCI OCPU at 2 vCPU per OCPU."],
                    confidence="high" if family else "medium",
                )
            )
        if line.aws_memory_gib > 0:
            terms = [family, "memory", "count"] if family else ["memory", "count"]
            preferred = [f"{family}-memory-count"] if family else []
            drafts.append(
                RequirementDraft(
                    service_name="compute",
                    metric="compute_memory",
                    quantity=_round(instance_count * line.aws_memory_gib),
                    unit="gb",
                    terms=terms,
                    preferred_names=preferred,
                    assumptions=assumptions + ["AWS GiB memory values were treated as GB for limit matching."],
                    confidence="high" if family else "medium",
                    avoid_terms=["reserved", "reservable", "free", "backup", "core"],
                )
            )
        return drafts

    def _is_database_line(self, line: BillComparisonLine) -> bool:
        text = line.search_text
        return bool(
            "postgresql service" in text
            or "postgresql.vm." in text
            or "aurora postgresql" in text
            or "mysql" in text
            or "relational database service" in text and "postgresql" in text
        )

    def _database_drafts(self, line: BillComparisonLine) -> list[RequirementDraft]:
        text = line.search_text
        if not line.is_hour_usage():
            return []
        count = self._steady_state_count(line)
        assumptions = [
            f"Database usage hours were divided by {self.hours_per_month:g} to estimate DB system count."
        ]
        if "postgresql" in text:
            family = _db_family(line)
            preferred = ["dbsystem-count"]
            terms = ["dbsystem", "count"]
            if family:
                preferred.insert(0, f"dbsystem-{family}-count")
                terms.insert(1, family)
            return [
                RequirementDraft(
                    service_name="postgresql",
                    metric="postgresql_dbsystem_count",
                    quantity=_round(count),
                    unit="dbsystem",
                    terms=terms,
                    preferred_names=preferred,
                    assumptions=assumptions,
                    confidence="high" if family else "medium",
                )
            ]
        if "mysql" in text:
            return [
                RequirementDraft(
                    service_name="mysql",
                    metric="mysql_instance_count",
                    quantity=_round(count),
                    unit="dbsystem",
                    terms=["mysql", "count"],
                    preferred_names=[],
                    assumptions=assumptions,
                    confidence="low",
                )
            ]
        return []

    def _is_block_storage_line(self, line: BillComparisonLine) -> bool:
        text = line.search_text
        return bool(
            line.oci_part_num_block_vol
            or "oci block volume" in text
            or "ebs:" in text
            or "gp2" in text
            or "gp3" in text
            or "provisioned storage" in text
        )

    def _block_storage_drafts(self, line: BillComparisonLine) -> list[RequirementDraft]:
        text = line.search_text
        if "snapshot" in text or "backup" in text or "iops" in text:
            return []
        if not _looks_like_gb_storage(line):
            return []
        return [
            RequirementDraft(
                service_name="block-storage",
                metric="block_volume_storage",
                quantity=_round(line.quantity),
                unit="gb",
                terms=["total", "storage", "gb"],
                preferred_names=["total-storage-gb"],
                assumptions=["GB-month storage quantities were treated as average provisioned GB."],
                confidence="high",
            )
        ]

    def _is_object_storage_line(self, line: BillComparisonLine) -> bool:
        text = line.search_text
        return bool(line.oci_part_num_obj_stor or "oci object storage" in text)

    def _object_storage_drafts(self, line: BillComparisonLine) -> list[RequirementDraft]:
        text = line.search_text
        if "request" in text or "objects listed" in text:
            return []
        if not _looks_like_gb_storage(line):
            return []
        return [
            RequirementDraft(
                service_name="object-storage",
                metric="object_storage_capacity",
                quantity=_round(line.quantity * 1024 * 1024 * 1024),
                unit="bytes",
                terms=["storage", "bytes"],
                preferred_names=["storage-bytes"],
                assumptions=["GB-month object storage quantities were converted to bytes."],
                confidence="high",
                avoid_terms=["bucket"],
            )
        ]

    def _load_balancer_drafts(self, line: BillComparisonLine) -> list[RequirementDraft]:
        text = line.search_text
        if "gateway" in text:
            return []
        if "network loadbalancer" in text or "network load balancer" in text:
            return self._single_count_draft(
                line=line,
                service_name="network-load-balancer-api",
                metric="network_load_balancer_count",
                terms=["flexible", "network", "load", "balancer", "count"],
                preferred_names=["max-nlb-flexible-count"],
                unit="load_balancer",
                assumptions=["Network load balancer-hours were converted to steady-state NLB count."],
            )
        return self._single_count_draft(
            line=line,
            service_name="load-balancer",
            metric="load_balancer_count",
            terms=["flexible", "load", "balancer", "count"],
            preferred_names=["lb-flexible-count"],
            unit="load_balancer",
            assumptions=["Application load balancer-hours were converted to steady-state LB count."],
        )

    def _single_count_draft(
        self,
        *,
        line: BillComparisonLine,
        service_name: str,
        metric: str,
        terms: list[str],
        preferred_names: list[str],
        unit: str,
        assumptions: list[str],
        confidence: str = "medium",
    ) -> list[RequirementDraft]:
        if service_name not in self.catalog.services:
            return []
        return [
            RequirementDraft(
                service_name=service_name,
                metric=metric,
                quantity=_round(self._steady_state_count(line)),
                unit=unit,
                terms=terms,
                preferred_names=preferred_names,
                assumptions=assumptions,
                confidence=confidence,
            )
        ]

    def _steady_state_count(self, line: BillComparisonLine) -> float:
        if line.is_hour_usage():
            return line.quantity / self.hours_per_month
        return line.quantity

    def _append_unmapped(self, unmapped: list[dict[str, Any]], line: BillComparisonLine, reason: str) -> None:
        if len(unmapped) >= self.unmapped_limit:
            return
        unmapped.append(
            {
                "row_number": line.row_number,
                "source_sheet": line.source_sheet,
                "aws_product": line.aws_product,
                "aws_description": line.aws_description,
                "oci_service_name": line.oci_service_name,
                "quantity": line.quantity,
                "reason": reason,
            }
        )

    def _aggregate_to_output(self, record: dict[str, Any]) -> dict[str, Any]:
        confidence_values = record.pop("confidence_values")
        confidence = "low"
        if confidence_values and all(value == "high" for value in confidence_values):
            confidence = "high"
        elif any(value in {"high", "medium"} for value in confidence_values):
            confidence = "medium"

        return {
            "service_name": record["service_name"],
            "service_description": record["service_description"],
            "limit_name": record["limit_name"],
            "limit_description": record["limit_description"],
            "scope_type": record["scope_type"],
            "metric": record["metric"],
            "required_quantity": _round(record["required_quantity"]),
            "unit": record["unit"],
            "confidence": confidence,
            "source_line_count": len(record["source_rows"]),
            "source_rows": sorted(record["source_rows"])[:50],
            "source_products": [name for name, _ in record["source_products"].most_common(10)],
            "source_oci_services": [
                name for name, _ in record["source_oci_services"].most_common(10)
            ],
            "assumptions": sorted(record["assumptions"]),
            "evidence": record["evidence"],
        }

    def _service_summary_to_output(self, item: dict[str, Any]) -> dict[str, Any]:
        return {
            "service_name": item["service_name"],
            "service_description": item["service_description"],
            "line_items": item["line_items"],
            "source_quantity": _round(item["source_quantity"]),
            "aws_products": [name for name, _ in item["aws_products"].most_common(8)],
            "oci_service_names": [name for name, _ in item["oci_service_names"].most_common(8)],
        }


def load_bill_comparison_lines(path: Path) -> list[BillComparisonLine]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _load_csv_lines(path)
    if suffix == ".xlsx":
        return _load_xlsx_lines(path)
    raise ValueError(f"Unsupported bill comparison input type: {path.suffix}")


def _load_csv_lines(path: Path) -> list[BillComparisonLine]:
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
        reader = csv.DictReader(handle)
        _validate_headers(reader.fieldnames or [], path.name)
        return [
            BillComparisonLine.from_mapping(index, path.name, row)
            for index, row in enumerate(reader, start=2)
            if any(value not in (None, "") for value in row.values())
        ]


def _load_xlsx_lines(path: Path) -> list[BillComparisonLine]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_names = ["AWS_CSV_Output"] + [
        name for name in workbook.sheetnames if name != "AWS_CSV_Output"
    ]
    for sheet_name in sheet_names:
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        header_row, headers = _find_header_row(worksheet)
        if not headers:
            continue
        _validate_headers(headers, sheet_name)
        lines: list[BillComparisonLine] = []
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if not any(value not in (None, "") for value in values):
                continue
            data = dict(zip(headers, values, strict=False))
            lines.append(BillComparisonLine.from_mapping(row_number, sheet_name, data))
        return lines
    raise ValueError("No bill comparison tab with AWS_CSV_Output-style headers was found.")


def _find_header_row(worksheet) -> tuple[int, list[str]]:
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(20, worksheet.max_row), values_only=True),
        start=1,
    ):
        headers = [_clean(value) for value in row]
        normalized = {header for header in headers if header}
        if REQUIRED_BILL_COLUMNS.issubset(normalized):
            return row_number, headers
    return 0, []


def _validate_headers(headers: Iterable[str], source: str) -> None:
    missing = REQUIRED_BILL_COLUMNS - {header for header in headers if header}
    if missing:
        raise ValueError(f"{source} is missing required bill comparison columns: {sorted(missing)}")


def _compute_family(line: BillComparisonLine) -> str | None:
    text = line.search_text
    explicit = _extract_vm_family(text)
    if explicit:
        return explicit
    if "a2.standard" in text or "standard a2" in text:
        return "standard-a2"
    if "standard3" in text or "x9" in text:
        return "standard3"
    if "optimized3" in text:
        return "optimized3"
    if "standard.e6" in text or "amd e6" in text:
        return "standard-e6"
    if "standard.e5" in text:
        return "standard-e5"
    if "standard.e4" in text:
        return "standard-e4"
    return None


def _db_family(line: BillComparisonLine) -> str | None:
    family = _compute_family(line)
    if family == "standard-e6":
        return "e6"
    if family == "standard-e5":
        return "e5"
    if family == "standard3":
        return "standard3"
    return None


def _extract_vm_family(text: str) -> str | None:
    normalized = text.lower()
    match = re.search(r"vm\.([a-z]+)\.([a-z0-9]+)\.flex", normalized)
    if match:
        return f"{match.group(1)}-{match.group(2)}"
    match = re.search(r"vm\.(standard[0-9]+)\.flex", normalized)
    if match:
        return match.group(1)
    match = re.search(r"vm\.(optimized[0-9]+)\.flex", normalized)
    if match:
        return match.group(1)
    return None


def _gpu_family(line: BillComparisonLine) -> str | None:
    text = line.search_text
    if "a10" in text:
        return "gpu-a10"
    if "l40s" in text:
        return "gpu-l40s"
    if "a100" in text:
        return "gpu-a100-v2"
    if "h100" in text:
        return "gpu-h100"
    if "b200" in text:
        return "gpu-b200"
    return None


def _looks_like_gb_storage(line: BillComparisonLine) -> bool:
    text = line.search_text
    if any(marker in text for marker in ["gb-month", "gb-month", "gb-mo", "gb / month", "gb-month"]):
        return True
    if "provisioned storage" in text or "storage used" in text or "volume" in text:
        return line.quantity > 0
    return False
