from __future__ import annotations

import csv
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.models import BomDocument, BomItem, BomRecommendation, LimitItem
from app.services.bill_compare import (
    BillComparisonAnalyzer,
    LimitCatalog,
    LimitDefinition,
    LimitService,
    load_bill_comparison_lines,
)


@dataclass
class ExtractedResource:
    resource_type: str
    display_name: str
    quantity: float
    unit: str
    confidence: str
    source_excerpt: str
    assumptions: list[str]


LIMIT_MAPPING = {
    "compute_instances": ("compute", ["instance-count", "standard", "vm", "count"]),
    "ocpus": ("compute", ["core-count", "ocpu", "cores"]),
    "memory_gb": ("compute", ["memory", "gb"]),
    "block_volume_gb": ("block-storage", ["total-storage-gb", "storage-gb"]),
    "block_volume_count": ("block-storage", ["volume-count"]),
    "object_storage_bucket": ("object-storage", ["bucket-count", "bucket"]),
    "load_balancer": ("load-balancer", ["lb-flexible-count", "lb-", "count"]),
    "vcn": ("vcn", ["vcn-count"]),
    "subnet": ("vcn", ["subnet-count"]),
    "public_ip": ("vcn", ["reserved-public-ip-count", "public-ip"]),
    "nat_gateway": ("vcn", ["nat-gateway-count"]),
    "drg": ("vcn", ["drg-count"]),
    "oke_cluster": ("container-engine", ["cluster-count"]),
    "file_storage": ("filesystem", ["file-system-count", "filesystem", "mount-target"]),
    "database": ("database", ["database", "db-system", "autonomous", "count"]),
}

RESOURCE_PATTERNS = [
    ("compute_instances", r"(\d+(?:\.\d+)?)\s+(?:compute\s+)?(?:instances?|vms?|virtual machines?)", "count"),
    ("ocpus", r"(\d+(?:\.\d+)?)\s*(?:ocpus?|oCPU|cores?)", "ocpu"),
    ("memory_gb", r"(\d+(?:\.\d+)?)\s*(?:gb|gib)\s+(?:memory|ram)", "gb"),
    ("block_volume_gb", r"(\d+(?:\.\d+)?)\s*(?:gb|gib|tb|tib)\s+(?:block\s+)?volumes?", "gb"),
    ("block_volume_count", r"(\d+(?:\.\d+)?)\s+(?:block\s+)?volumes?", "count"),
    ("object_storage_bucket", r"(\d+(?:\.\d+)?)\s+(?:object\s+storage\s+)?buckets?", "count"),
    ("load_balancer", r"(\d+(?:\.\d+)?)\s+(?:load\s+balancers?|lbs?)", "count"),
    ("vcn", r"(\d+(?:\.\d+)?)\s+(?:vcns?|virtual cloud networks?)", "count"),
    ("subnet", r"(\d+(?:\.\d+)?)\s+subnets?", "count"),
    ("public_ip", r"(\d+(?:\.\d+)?)\s+(?:public|reserved public)\s+ips?", "count"),
    ("nat_gateway", r"(\d+(?:\.\d+)?)\s+nat\s+gateways?", "count"),
    ("drg", r"(\d+(?:\.\d+)?)\s+(?:drgs?|dynamic routing gateways?)", "count"),
    ("oke_cluster", r"(\d+(?:\.\d+)?)\s+(?:oke|kubernetes|container engine)\s+clusters?", "count"),
    ("file_storage", r"(\d+(?:\.\d+)?)\s+(?:file systems?|mount targets?)", "count"),
    ("database", r"(\d+(?:\.\d+)?)\s+(?:databases?|db systems?|autonomous databases?)", "count"),
]


class BomAnalyzer:
    def __init__(self, db: Session, settings: Settings) -> None:
        self.db = db
        self.settings = settings

    def analyze_upload(self, filename: str, content_type: str | None, file: BinaryIO) -> BomDocument:
        content = file.read()
        if len(content) > self.settings.upload_max_bytes:
            raise ValueError("Upload exceeds configured maximum size.")

        self.settings.upload_dir.mkdir(parents=True, exist_ok=True)
        digest = hashlib.sha256(content).hexdigest()
        suffix = Path(filename).suffix.lower()
        stored_path = self.settings.upload_dir / f"{digest}{suffix}"
        stored_path.write_bytes(content)

        text = self.extract_text(stored_path, content_type)
        doc = BomDocument(
            filename=filename,
            content_type=content_type,
            size_bytes=len(content),
            sha256=digest,
            status="analyzed",
            raw_text_preview=text[:4000],
        )
        self.db.add(doc)
        self.db.flush()

        if self._analyze_bill_comparison_upload(doc.id, stored_path, filename):
            self.db.commit()
            self.db.refresh(doc)
            return doc

        resources = self.extract_resources(text)
        for resource in resources:
            item = BomItem(
                document_id=doc.id,
                resource_type=resource.resource_type,
                display_name=resource.display_name,
                quantity=resource.quantity,
                unit=resource.unit,
                confidence=resource.confidence,
                source_excerpt=resource.source_excerpt,
                assumptions=resource.assumptions,
                service_name=LIMIT_MAPPING.get(resource.resource_type, (None, []))[0],
            )
            self.db.add(item)
            self.db.flush()
            self._create_recommendation(doc.id, item)

        if not resources:
            self.db.add(
                BomRecommendation(
                    document_id=doc.id,
                    required_quantity=0,
                    confidence="low",
                    explanation=(
                        "No OCI resource quantities were detected. Upload Terraform plan JSON "
                        "or include explicit quantities such as '3 load balancers'."
                    ),
                    assumptions=["Manual review required."],
                )
            )

        self.db.commit()
        self.db.refresh(doc)
        return doc

    def _analyze_bill_comparison_upload(
        self, document_id: str, path: Path, filename: str
    ) -> bool:
        if path.suffix.lower() not in {".csv", ".xlsx"}:
            return False
        try:
            lines = load_bill_comparison_lines(path)
        except ValueError:
            return False

        if not lines:
            self.db.add(
                BomRecommendation(
                    document_id=document_id,
                    required_quantity=0,
                    confidence="low",
                    explanation="Bill comparison workbook was detected, but it did not contain bill rows.",
                    assumptions=["Manual review required."],
                )
            )
            return True

        catalog = self._live_limit_catalog()
        if not catalog.services:
            self.db.add(
                BomRecommendation(
                    document_id=document_id,
                    required_quantity=0,
                    confidence="low",
                    explanation=(
                        "Bill comparison workbook was detected, but no persisted numeric OCI limits "
                        "are available for comparison yet."
                    ),
                    assumptions=["Run a limit scan before relying on bill comparison recommendations."],
                )
            )
            return True

        analysis = BillComparisonAnalyzer(catalog).analyze_lines(lines, source_file=filename)
        requirements = analysis.get("requirements", [])
        for requirement in requirements:
            item = BomItem(
                document_id=document_id,
                resource_type=requirement["metric"],
                display_name=requirement["metric"].replace("_", " ").title(),
                quantity=requirement["required_quantity"],
                unit=requirement["unit"],
                region=None,
                service_name=requirement["service_name"],
                confidence=requirement["confidence"],
                source_excerpt="\n".join(requirement.get("evidence", [])[:4])[:2000],
                assumptions=self._bill_requirement_assumptions(requirement),
            )
            self.db.add(item)
            self.db.flush()
            self._create_bill_requirement_recommendation(document_id, item, requirement)

        summary = analysis.get("summary", {})
        unmapped_count = int(summary.get("unmapped_line_items") or 0)
        if not requirements:
            self.db.add(
                BomRecommendation(
                    document_id=document_id,
                    required_quantity=0,
                    confidence="low",
                    explanation=(
                        "Bill comparison workbook was parsed, but no rows could be mapped to "
                        "persisted OCI service limits."
                    ),
                    assumptions=self._bill_summary_assumptions(analysis),
                )
            )
        elif unmapped_count:
            self.db.add(
                BomRecommendation(
                    document_id=document_id,
                    required_quantity=0,
                    confidence="low",
                    explanation=(
                        f"{unmapped_count} bill comparison row(s) were not mapped to OCI limits "
                        "and need manual review."
                    ),
                    assumptions=self._bill_summary_assumptions(analysis),
                )
            )
        return True

    def _live_limit_catalog(self) -> LimitCatalog:
        items = list(
            self.db.scalars(
                select(LimitItem)
                .where(LimitItem.last_allowed_limit.is_not(None))
                .order_by(LimitItem.service_name, LimitItem.limit_name)
            )
        )
        services = {
            item.service_name: LimitService(item.service_name, item.service_name)
            for item in items
            if item.service_name
        }
        definitions = [
            LimitDefinition(
                service_name=item.service_name,
                name=item.limit_name,
                description=item.resource_name or item.limit_name,
                scope_type=item.scope_type,
                is_resource_availability_supported=item.last_collection_status == "ok",
            )
            for item in items
            if item.service_name and item.limit_name
        ]
        return LimitCatalog(services=services.values(), definitions=definitions)

    def _create_bill_requirement_recommendation(
        self, document_id: str, item: BomItem, requirement: dict
    ) -> None:
        match = self._match_bill_requirement(requirement)
        if not match:
            self.db.add(
                BomRecommendation(
                    document_id=document_id,
                    bom_item_id=item.id,
                    matching_service=requirement["service_name"],
                    matching_limit_name=requirement["limit_name"],
                    required_quantity=item.quantity,
                    confidence="low",
                    explanation=(
                        f"Mapped the bill row to {requirement['service_name']}/"
                        f"{requirement['limit_name']}, but that exact numeric limit is not present "
                        "in persisted LIP scan data."
                    ),
                    assumptions=item.assumptions,
                )
            )
            return

        self._add_capacity_recommendation(
            document_id=document_id,
            item=item,
            match=match,
            base_confidence=item.confidence,
            explanation_prefix=(
                f"Bill comparison mapping matched {item.display_name} to "
                f"{match.service_name}/{match.limit_name}. "
            ),
        )

    def _match_bill_requirement(self, requirement: dict) -> LimitItem | None:
        candidates = list(
            self.db.scalars(
                select(LimitItem)
                .where(LimitItem.service_name == requirement["service_name"])
                .where(LimitItem.limit_name == requirement["limit_name"])
                .where(LimitItem.last_allowed_limit.is_not(None))
            )
        )
        if not candidates:
            return None
        candidates.sort(
            key=lambda item: (
                item.last_available if item.last_available is not None else float("inf"),
                -(item.last_percent_used or 0),
                item.availability_domain or "",
            )
        )
        return candidates[0]

    def _bill_requirement_assumptions(self, requirement: dict) -> list[str]:
        assumptions = list(requirement.get("assumptions") or [])
        source_rows = requirement.get("source_rows") or []
        source_line_count = requirement.get("source_line_count")
        if source_line_count:
            assumptions.append(f"Aggregated from {source_line_count} bill comparison row(s).")
        if source_rows:
            assumptions.append(f"Source rows: {', '.join(str(row) for row in source_rows[:12])}.")
        return assumptions

    def _bill_summary_assumptions(self, analysis: dict) -> list[str]:
        summary = analysis.get("summary", {})
        assumptions = [
            f"Parsed {summary.get('line_items', 0)} bill comparison line item(s).",
            f"Mapped {summary.get('mapped_line_items', 0)} line item(s).",
        ]
        reason_counts = summary.get("unmapped_reason_counts") or {}
        if reason_counts:
            assumptions.append(f"Unmapped reasons: {json.dumps(reason_counts, sort_keys=True)}.")
        return assumptions

    def extract_text(self, path: Path, content_type: str | None) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if suffix == ".docx":
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        if suffix == ".xlsx":
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    rows.append(" ".join("" if cell is None else str(cell) for cell in row))
            return "\n".join(rows)
        if suffix == ".csv":
            with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
                reader = csv.reader(handle)
                return "\n".join(" ".join(row) for row in reader)
        if suffix in {".json", ".tfplan"}:
            return self._flatten_json(path)
        return path.read_text(encoding="utf-8", errors="ignore")

    def _flatten_json(self, path: Path) -> str:
        try:
            payload = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
        except json.JSONDecodeError:
            return path.read_text(encoding="utf-8", errors="ignore")
        text_parts: list[str] = []

        def walk(obj):
            if isinstance(obj, dict):
                if "type" in obj:
                    text_parts.append(str(obj["type"]))
                if "name" in obj:
                    text_parts.append(str(obj["name"]))
                for value in obj.values():
                    walk(value)
            elif isinstance(obj, list):
                for item in obj:
                    walk(item)
            else:
                text_parts.append(str(obj))

        walk(payload)
        return "\n".join(text_parts)

    def extract_resources(self, text: str) -> list[ExtractedResource]:
        resources: list[ExtractedResource] = []
        normalized = re.sub(r"\s+", " ", text)

        for resource_type, pattern, unit in RESOURCE_PATTERNS:
            for match in re.finditer(pattern, normalized, flags=re.IGNORECASE):
                quantity = float(match.group(1))
                excerpt = normalized[max(0, match.start() - 80) : match.end() + 80]
                if unit == "gb" and re.search(r"\b(?:tb|tib)\b", match.group(0), flags=re.IGNORECASE):
                    quantity *= 1024
                resources.append(
                    ExtractedResource(
                        resource_type=resource_type,
                        display_name=resource_type.replace("_", " ").title(),
                        quantity=quantity,
                        unit=unit,
                        confidence="medium",
                        source_excerpt=excerpt,
                        assumptions=["Detected by deterministic resource pattern matching."],
                    )
                )

        terraform_counts = self._extract_terraform_resource_counts(text)
        for resource_type, quantity in terraform_counts.items():
            resources.append(
                ExtractedResource(
                    resource_type=resource_type,
                    display_name=resource_type.replace("_", " ").title(),
                    quantity=quantity,
                    unit="count",
                    confidence="high",
                    source_excerpt="Terraform resource type detected.",
                    assumptions=["Quantity is count of matching Terraform resource declarations."],
                )
            )

        return self._merge_resources(resources)

    def _extract_terraform_resource_counts(self, text: str) -> dict[str, float]:
        mapping = {
            "oci_core_instance": "compute_instances",
            "oci_core_volume": "block_volume_count",
            "oci_objectstorage_bucket": "object_storage_bucket",
            "oci_load_balancer_load_balancer": "load_balancer",
            "oci_core_vcn": "vcn",
            "oci_core_subnet": "subnet",
            "oci_core_public_ip": "public_ip",
            "oci_core_nat_gateway": "nat_gateway",
            "oci_core_drg": "drg",
            "oci_containerengine_cluster": "oke_cluster",
            "oci_file_storage_file_system": "file_storage",
            "oci_database_db_system": "database",
            "oci_database_autonomous_database": "database",
        }
        counts: dict[str, float] = {}
        for terraform_type, resource_type in mapping.items():
            count = len(re.findall(terraform_type, text, flags=re.IGNORECASE))
            if count:
                counts[resource_type] = counts.get(resource_type, 0) + count
        return counts

    def _merge_resources(self, resources: list[ExtractedResource]) -> list[ExtractedResource]:
        merged: dict[tuple[str, str], ExtractedResource] = {}
        for resource in resources:
            key = (resource.resource_type, resource.unit)
            if key in merged:
                merged[key].quantity += resource.quantity
                merged[key].source_excerpt = f"{merged[key].source_excerpt}\n{resource.source_excerpt}"[:2000]
            else:
                merged[key] = resource
        return list(merged.values())

    def _create_recommendation(self, document_id: str, item: BomItem) -> None:
        match = self._match_limit(item)
        if not match:
            self.db.add(
                BomRecommendation(
                    document_id=document_id,
                    bom_item_id=item.id,
                    required_quantity=item.quantity,
                    confidence="low",
                    explanation=(
                        f"{item.display_name} could not be mapped to a scanned OCI service limit. "
                        "Manual review is required."
                    ),
                    assumptions=item.assumptions,
                )
            )
            return

        self._add_capacity_recommendation(
            document_id=document_id,
            item=item,
            match=match,
            base_confidence=item.confidence,
            explanation_prefix=(
                f"Mapped {item.display_name} to {match.service_name}/{match.limit_name}. "
            ),
        )

    def _add_capacity_recommendation(
        self,
        *,
        document_id: str,
        item: BomItem,
        match: LimitItem,
        base_confidence: str,
        explanation_prefix: str,
    ) -> None:
        current_usage = match.last_used or 0
        allowed = match.last_allowed_limit
        available = match.last_available
        available_capacity = available
        needed = False
        recommended = None
        explanation = explanation_prefix + "Current persisted scan data was used for the recommendation."
        if allowed is not None:
            projected_usage = current_usage + item.quantity
            available_capacity = available if available is not None else max(0, allowed - current_usage)
            needed = item.quantity > available_capacity or projected_usage > allowed * 0.9
            if needed:
                recommended = max(allowed + item.quantity, projected_usage / 0.8)
            else:
                explanation += " Available capacity appears sufficient."
        else:
            explanation += " The matched limit has no numeric allowed value; manual validation is needed."

        self.db.add(
            BomRecommendation(
                document_id=document_id,
                bom_item_id=item.id,
                limit_item_id=match.id,
                matching_service=match.service_name,
                matching_limit_name=match.limit_name,
                current_usage=current_usage,
                allowed_limit=allowed,
                available_capacity=available_capacity,
                required_quantity=item.quantity,
                limit_increase_needed=needed,
                recommended_new_limit=round(recommended, 2) if recommended is not None else None,
                confidence="high" if base_confidence == "high" else "medium",
                explanation=explanation,
                assumptions=item.assumptions,
            )
        )

    def _match_limit(self, item: BomItem) -> LimitItem | None:
        service_name, keywords = LIMIT_MAPPING.get(item.resource_type, (None, []))
        if not service_name:
            return None
        candidates = list(
            self.db.scalars(
                select(LimitItem)
                .where(LimitItem.service_name == service_name)
                .where(LimitItem.last_allowed_limit.is_not(None))
            )
        )
        if not candidates:
            return None
        scored: list[tuple[int, LimitItem]] = []
        for candidate in candidates:
            name = candidate.limit_name.lower()
            score = sum(1 for keyword in keywords if keyword.lower() in name)
            if item.unit in {"gb", "ocpu"} and item.unit in name:
                score += 2
            scored.append((score, candidate))
        scored.sort(key=lambda pair: (pair[0], pair[1].last_allowed_limit or 0), reverse=True)
        return scored[0][1] if scored and scored[0][0] > 0 else None
